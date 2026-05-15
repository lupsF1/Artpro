"""知识库文档解析与入库。"""

from __future__ import annotations

import json
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import KbChunk, KbDocument
from app.services.embeddings import EmbeddingUnavailableError, EmbeddingUpstreamError
from app.services.embeddings import embed_texts
from app.services.kb_chunking import chunk_text, chunk_text_with_page_bounds

_API_ROOT = Path(__file__).resolve().parents[2]
_SECTION_RE = re.compile(r"^[一二三四五六七八九十]+[、.．]\s*.+")
_SUBSECTION_RE = re.compile(r"^（[一二三四五六七八九十]+）\s*.+")
_DIRECTION_RE = re.compile(r"^(?:研究方向[:：]\s*)?\d{2}[\u4e00-\u9fffA-Za-z].*")
_SUBJECT_RE = re.compile(r"^\d{3,4}[\u4e00-\u9fffA-Za-z].*")
_OUTLINE_RE = re.compile(r"^《[^》]+》考试大纲$")


def _safe_filename(name: str, max_len: int = 180) -> str:
    base = (name or "upload").strip()
    base = re.sub(r"[^\w.\-()\u4e00-\u9fff]+", "_", base)
    return base[:max_len] if len(base) > max_len else base


def _parse_pdf(raw: bytes) -> list[tuple[str, dict]]:
    reader = PdfReader(BytesIO(raw))
    page_texts: list[str] = []
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        page_texts.append(t)
    outline_lines: list[str] = []
    for page_text in page_texts:
        outline_lines.extend([line.strip() for line in page_text.splitlines() if line.strip()])
    outline_pairs = _parse_exam_outline_lines(outline_lines, source="pdf")
    if outline_pairs:
        return outline_pairs
    return chunk_text_with_page_bounds(page_texts)


def _approx_tokens(text: str) -> int:
    cjk = len(re.findall(r"[\u4e00-\u9fff]", text))
    latin_words = len(re.findall(r"[A-Za-z0-9]+", text))
    punctuation = len(re.findall(r"[^\s\w\u4e00-\u9fff]", text))
    return cjk + latin_words + max(1, punctuation // 3)


def _breadcrumb_meta(state: dict[str, Any], *, source: str, block_type: str) -> dict[str, Any]:
    directions = list(state.get("researchDirections") or [])
    subjects = list(state.get("examSubjects") or [])
    return {
        "source": source,
        "blockType": block_type,
        "degreeType": state.get("degreeType"),
        "researchDirections": directions,
        "examSubjects": subjects,
        "outlineName": state.get("outlineName"),
        "sectionTitle": state.get("sectionTitle"),
        "subsectionTitle": state.get("subsectionTitle"),
        "breadcrumb": {
            "degreeType": state.get("degreeType"),
            "researchDirections": directions,
            "examSubjects": subjects,
            "outlineName": state.get("outlineName"),
            "sectionTitle": state.get("sectionTitle"),
            "subsectionTitle": state.get("subsectionTitle"),
        },
    }


def _split_outline_block(
    lines: list[str],
    meta: dict[str, Any],
    *,
    min_tokens: int = 300,
    max_tokens: int = 800,
    overlap_ratio: float = 0.1,
) -> list[tuple[str, dict]]:
    text = "\n".join(lines).strip()
    if not text:
        return []
    if _approx_tokens(text) <= max_tokens:
        return [(text, {**meta, "partIndex": 0, "tokenApprox": _approx_tokens(text)})]

    chunks: list[tuple[str, dict]] = []
    current: list[str] = []
    current_tokens = 0
    overlap_tokens = max(1, int(max_tokens * overlap_ratio))

    def flush() -> None:
        nonlocal current, current_tokens
        if not current:
            return
        chunk_text_value = "\n".join(current).strip()
        if chunk_text_value:
            chunks.append(
                (
                    chunk_text_value,
                    {
                        **meta,
                        "partIndex": len(chunks),
                        "tokenApprox": _approx_tokens(chunk_text_value),
                        "overlapRatio": overlap_ratio,
                    },
                )
            )
        overlap: list[str] = []
        total = 0
        for line in reversed(current):
            line_tokens = _approx_tokens(line)
            if overlap and total + line_tokens > overlap_tokens:
                break
            overlap.insert(0, line)
            total += line_tokens
        current = overlap
        current_tokens = total

    for line in lines:
        line_tokens = _approx_tokens(line)
        if current and current_tokens + line_tokens > max_tokens:
            flush()
        current.append(line)
        current_tokens += line_tokens
        if current_tokens >= min_tokens and current_tokens >= max_tokens * 0.8:
            flush()
    if current:
        chunk_text_value = "\n".join(current).strip()
        if not chunks or chunk_text_value != chunks[-1][0]:
            chunks.append(
                (
                    chunk_text_value,
                    {
                        **meta,
                        "partIndex": len(chunks),
                        "tokenApprox": _approx_tokens(chunk_text_value),
                        "overlapRatio": overlap_ratio,
                    },
                )
            )
    return chunks


def _build_meta_prefix(state: dict[str, Any]) -> str:
    parts: list[str] = []
    if state.get("degreeType"):
        parts.append(state["degreeType"])
    dirs = state.get("researchDirections") or []
    if dirs:
        parts.append("研究方向：" + " ".join(dirs))
    subs = state.get("examSubjects") or []
    if subs:
        parts.append("考试科目：" + " ".join(subs))
    if state.get("outlineName"):
        parts.append(state["outlineName"])
    return " ".join(parts)


def _parse_exam_outline_lines(lines: list[str], *, source: str) -> list[tuple[str, dict]]:
    state: dict[str, Any] = {
        "degreeType": None,
        "researchDirections": [],
        "examSubjects": [],
        "outlineName": None,
        "sectionTitle": None,
        "subsectionTitle": None,
    }
    out: list[tuple[str, dict]] = []
    current_lines: list[str] = []
    current_meta: dict[str, Any] | None = None
    saw_outline_marker = False

    def flush() -> None:
        nonlocal current_lines, current_meta
        if current_meta and current_lines:
            prefix = _build_meta_prefix(state)
            lines_with_meta = ([prefix] + current_lines) if prefix else current_lines
            out.extend(_split_outline_block(lines_with_meta, current_meta))
        current_lines = []
        current_meta = None

    for raw in lines:
        line = " ".join((raw or "").split())
        if not line:
            continue

        if line in {"学术学位", "专业学位"} or line.endswith("学位"):
            state["degreeType"] = line
            state["researchDirections"] = []
            state["examSubjects"] = []
            saw_outline_marker = True

        if line.startswith("研究方向"):
            value = line.split("：", 1)[-1].split(":", 1)[-1].strip()
            if value:
                state["researchDirections"].append(value)
            saw_outline_marker = True
        elif _DIRECTION_RE.match(line) and state.get("degreeType") and not state.get("examSubjects"):
            state["researchDirections"].append(line)
            saw_outline_marker = True

        if line.startswith("考试科目"):
            value = line.split("：", 1)[-1].split(":", 1)[-1].strip()
            if value:
                state["examSubjects"].append(value)
            saw_outline_marker = True
        elif _SUBJECT_RE.match(line) and state.get("examSubjects") is not None and state.get("researchDirections"):
            state["examSubjects"].append(line)
            saw_outline_marker = True

        if _OUTLINE_RE.match(line):
            flush()
            state["outlineName"] = line
            state["sectionTitle"] = None
            state["subsectionTitle"] = None
            saw_outline_marker = True
            continue

        if _SECTION_RE.match(line):
            if current_meta is None:
                # Starting a new top-level section
                state["sectionTitle"] = line
                state["subsectionTitle"] = None
                saw_outline_marker = True
                current_meta = _breadcrumb_meta(state, source=source, block_type="section")
                current_lines = [line]
                continue
            else:
                # Already inside a section - treat as subsection content
                state["subsectionTitle"] = line
                saw_outline_marker = True

        if _SUBSECTION_RE.match(line):
            state["subsectionTitle"] = line
            saw_outline_marker = True

        if current_meta is None:
            continue
        current_lines.append(line)

    flush()
    return out if saw_outline_marker and out else []


def _parse_docx(raw: bytes) -> list[tuple[str, dict]]:
    doc = Document(BytesIO(raw))
    paras: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            paras.append(t)
    outline_pairs = _parse_exam_outline_lines(paras, source="docx")
    if outline_pairs:
        return outline_pairs
    text = "\n\n".join(paras)
    pairs = chunk_text(text)
    return [(c, {**m, "source": "docx"}) for c, m in pairs]


async def ingest_file_to_document(
    db: AsyncSession,
    *,
    raw: bytes,
    filename: str,
    title_override: str | None,
) -> KbDocument:
    fn = filename.lower()
    mime: str | None = None
    if fn.endswith(".pdf"):
        mime = "application/pdf"
    elif fn.endswith(".docx"):
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    title = (title_override or "").strip() or Path(filename).stem or "未命名文档"
    doc_row = KbDocument(
        title=title[:500],
        source_type="upload",
        source_filename=filename[:500],
        mime_type=mime,
        status="processing",
        review_status="draft",
    )
    db.add(doc_row)
    await db.flush()

    upload_dir = _API_ROOT / "data" / "kb_uploads"
    upload_dir.mkdir(parents=True, exist_ok=True)
    spath = upload_dir / f"{doc_row.id}_{_safe_filename(filename)}"
    try:
        spath.write_bytes(raw)
        doc_row.storage_path = str(spath.relative_to(_API_ROOT))
    except OSError:
        doc_row.storage_path = None

    await db.commit()
    await db.refresh(doc_row)

    try:
        if fn.endswith(".pdf"):
            pairs = _parse_pdf(raw)
        elif fn.endswith(".docx"):
            pairs = _parse_docx(raw)
        else:
            doc_row.status = "failed"
            doc_row.error_message = "仅支持 .pdf 或 .docx"
            await db.commit()
            await db.refresh(doc_row)
            return doc_row

        if not pairs:
            doc_row.status = "failed"
            doc_row.error_message = "未能从文档中提取文本（扫描件需 OCR，见文档说明）"
            await db.commit()
            await db.refresh(doc_row)
            return doc_row

        texts = [p[0] for p in pairs]
        batch_size = 16
        all_vecs: list[list[float]] = []
        for i in range(0, len(texts), batch_size):
            batch = texts[i : i + batch_size]
            vecs = await embed_texts(batch)
            all_vecs.extend(vecs)

        for ord_i, ((ctext, meta), vec) in enumerate(zip(pairs, all_vecs, strict=True)):
            ch = KbChunk(
                document_id=doc_row.id,
                ordinal=ord_i,
                text=ctext,
                meta_json=json.dumps(meta, ensure_ascii=False),
                embedding_json=json.dumps(vec, ensure_ascii=False),
            )
            db.add(ch)

        doc_row.status = "ready"
        doc_row.error_message = None
        await db.commit()
        await db.refresh(doc_row)
        return doc_row
    except (EmbeddingUnavailableError, EmbeddingUpstreamError) as e:
        doc_row.status = "failed"
        doc_row.error_message = str(e)[:2000]
        await db.commit()
        await db.refresh(doc_row)
        return doc_row
    except Exception as e:
        doc_row.status = "failed"
        doc_row.error_message = str(e)[:2000]
        await db.commit()
        await db.refresh(doc_row)
        return doc_row
